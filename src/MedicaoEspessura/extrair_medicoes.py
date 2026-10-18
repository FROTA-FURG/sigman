#!/usr/bin/env python3
"""
Extrai as medições de espessura da planilha "Controle de espessura.xlsx"
(ferramenta de gestão do monitoramento de chapas) para o JSON que a tela
Medição de Espessura usa:

    resources/js/Pages/HullThickness/data/medicoes.json

Uso (da raiz do projeto):
    src/.venv/bin/python src/MedicaoEspessura/extrair_medicoes.py [caminho.xlsx]

Estrutura da planilha: uma aba por embarcação+região (AS_BE, AS_BB, AS_CP,
AS_CC, AS_AP, CMI_...). Cada chapa é um bloco:

    CHAPA 34 BE \\n COSTADO
    N°    1    2    3    4    5    LI
    E.P   7.94 7.94 7.94 7.94 7.94 6.352      <- espessura de projeto por ponto; LI = 80% da E.P.
    2018  7.3  8.1  ...                       <- uma linha por campanha (ano)
    2022  ...

As abas têm 49 blocos de modelo; só os que têm E.P. e alguma medição entram.
A aba _AP junta antepara de vante, antepara de ré e espelho de popa: aqui ela
é separada nas três regiões do desenho, renumerando cada uma a partir de 1.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
PLANILHA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'docs' / 'Controle de espessura.xlsx'
SAIDA = RAIZ / 'resources/js/Pages/HullThickness/data/medicoes.json'

EMBARCACOES = {
    'AS': {'construcao': 1973},
    'CMI': {'construcao': 2017},
}
# Aba _AP -> regiões do desenho. No AS a categoria do bloco diz a região; no
# CMI os rótulos da aba vieram do modelo do AS (3 AV + 3 RE), então a divisão
# segue a quantidade de chapas do desenho do CMI: 2 AV, 3 AR, 1 espelho.
CATEGORIA_AP = {'ANTEPARA AV': 'AV', 'ANTEPARA RE': 'AR', 'ESPELHO POPA': 'EP'}
DIVISAO_AP_CMI = ['AV', 'AV', 'AR', 'AR', 'AR', 'EP']
# Anos digitados errado na planilha (AS_AP chapa 2: "6" na linha de 2024).
CORRECOES_ANO = {6: 2024}


def numero(v):
    return v if isinstance(v, (int, float)) else None


def ler_blocos(ws):
    blocos = []
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.strip().upper().startswith('CHAPA'):
                blocos.append((cell.row, cell.column, v.strip()))
    for r, c, titulo in blocos:
        m = re.match(r'CHAPA\s+(\d+)\s*[A-Z]*\s*(.*)', titulo, re.I | re.S)
        colunas = []
        k = c + 1
        while ws.cell(r + 1, k).value not in ('LI', None):
            colunas.append(k)
            k += 1
        eps = [numero(ws.cell(r + 2, col).value) for col in colunas]
        campanhas = {}
        linha = r + 3
        while numero(ws.cell(linha, c).value) is not None:
            ano = int(ws.cell(linha, c).value)
            valores = [numero(ws.cell(linha, col).value) for col in colunas]
            if any(x is not None for x in valores):
                campanhas[ano] = valores
            linha += 1
        if not campanhas or not any(eps):
            continue  # bloco de modelo, sem dados
        yield {
            'numero': int(m.group(1)),
            'categoria': ' '.join(m.group(2).split()).upper(),
            'ep': max(x for x in eps if x is not None),
            'campanhas': campanhas,
        }


def main():
    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    saida = {}
    observacoes = {emb: [] for emb in EMBARCACOES}

    for emb, cfg in EMBARCACOES.items():
        chapas = {}
        anos = set()
        for ws in wb.worksheets:
            if not ws.title.startswith(f'{emb}_'):
                continue
            sigla = ws.title.split('_', 1)[1]
            blocos = list(ler_blocos(ws))
            contador = {}
            for i, b in enumerate(blocos):
                if sigla == 'AP':
                    regiao = DIVISAO_AP_CMI[i] if emb == 'CMI' else CATEGORIA_AP.get(b['categoria'], 'AP')
                    contador[regiao] = contador.get(regiao, 0) + 1
                    n = contador[regiao]
                else:
                    regiao, n = sigla, b['numero']

                campanhas = {}
                for ano, valores in b['campanhas'].items():
                    if ano < 1900:
                        corrigido = CORRECOES_ANO.get(ano)
                        if corrigido is None:
                            observacoes[emb].append(f'{ws.title} chapa {b["numero"]}: ano "{ano}" ignorado (não reconhecido).')
                            continue
                        observacoes[emb].append(f'{ws.title} chapa {b["numero"]}: ano "{ano}" lido como {corrigido}.')
                        ano = corrigido
                    campanhas[str(ano)] = valores
                    anos.add(ano)
                chapas[f'{regiao}-{n}'] = {
                    'ep': b['ep'],
                    'tipo': b['categoria'].title() if sigla != 'AP' else None,
                    'origem': f'{ws.title} · chapa {b["numero"]}',
                    'medicoes': campanhas,
                }

        # Campanhas repetidas (mesmos valores em anos seguidos): avisa, não altera.
        lista_anos = sorted(anos)
        for a, b in zip(lista_anos, lista_anos[1:]):
            ambos = [c for c in chapas.values() if str(a) in c['medicoes'] and str(b) in c['medicoes']]
            iguais = sum(1 for c in ambos if c['medicoes'][str(a)] == c['medicoes'][str(b)])
            if ambos and iguais / len(ambos) > 0.5:
                observacoes[emb].append(
                    f'Campanha {b} idêntica à de {a} em {iguais} de {len(ambos)} chapas — parece cópia, não nova medição. '
                    'Isso achata a tendência (TC tendência menor que a real).'
                )
        acima = sum(1 for c in chapas.values() for v in c['medicoes'].values() for x in v if x is not None and x > c['ep'])
        total = sum(1 for c in chapas.values() for v in c['medicoes'].values() for x in v if x is not None)
        if total and acima / total > 0.3:
            observacoes[emb].append(
                f'{acima} de {total} leituras estão acima da espessura de projeto (E.P.). Vale confirmar a E.P. dessas chapas '
                '(tolerância de laminação, chapa substituída por outra mais grossa ou medição sobre pintura).'
            )
        if emb == 'CMI':
            observacoes[emb].append('Aba CMI_AP: dividida pelo desenho (antepara de vante 1–2, antepara de ré 1–3, espelho de popa 1) — confirmar.')
        if emb == 'AS':
            observacoes[emb].append('Aba AS_AP: chapas 1–3 = antepara de vante, 4–7 = antepara de ré (1–4), 8–10 = espelho de popa (1–3).')

        saida[emb] = {
            'construcao': cfg['construcao'],
            'campanhas': [{'id': str(a), 'data': f'{a}-07-01'} for a in lista_anos],
            'chapas': dict(sorted(chapas.items(), key=lambda kv: (kv[0].split('-')[0], int(kv[0].split('-')[1])))),
            'observacoes': observacoes[emb],
            'fonte': PLANILHA.name,
        }
        print(f'{emb}: {len(chapas)} chapas, campanhas {lista_anos}')
        for o in observacoes[emb]:
            print(f'   ! {o}')

    SAIDA.write_text(json.dumps(saida, ensure_ascii=False, indent=1))
    print(f'Gravado {SAIDA.relative_to(RAIZ)}')


if __name__ == '__main__':
    main()
